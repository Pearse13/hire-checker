"""
Hire report checker - ACS in-house hire report vs supplier hire reports.

Usage:
    python3 hire_check.py INHOUSE.xlsx SUPPLIER1.xlsx [SUPPLIER2.xlsx ...] [--code-col C]

For each supplier report it writes SUPPLIER_checked.xlsx where:
  - lines whose plant code is on any sheet of the in-house report -> green + bold,
    with the ACS sheet and cell they matched in an added "ACS match" column
  - lines NOT on the in-house report are left untouched and listed on screen
It also writes INHOUSE_checked.xlsx: ACS lines charged on the supplier reports -> bold + thick border,
keeping the report's own colours.
Once every report is checked it prints suggestions: what each unknown plant code probably is on the
ACS report (by date, rate, quantity and description), then where the ACS report's unpaired £0 items are.
The web app (src/engine.js) follows the same rules.
"""
import sys, os, re, subprocess, tempfile, warnings
from datetime import date, datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border, Side
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter

warnings.filterwarnings("ignore")

GREEN = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
THICK = Side(style="thick", color="000000")
NOT_CODES = {"CODE", "PLANTCODE", "PLANT", "ITEM", "ITEMCODE", "PRODUCT", "PRODUCTCODE",
             "REF", "REFERENCE", "STOCK", "STOCKCODE", "DESCRIPTION", "QTY", "QUANTITY", "DAYS", "RATE", "CHARGE", "AMOUNT", "TOTAL", "SUBTOTAL", "VAT", "NET", "GROSS"}
STOP = {"the", "and", "with", "for", "inc", "incl", "including", "donotuse", "set", "per", "week", "day",
        "standard", "heavy", "duty", "unit", "type", "new", "only", "from", "use"}
SPLIT = r"[+,;&]"


def norm(v):
    """Tidy a cell into a comparable code: uppercase, no spaces."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", "", str(v)).upper()


def looks_like_code(v):
    """A plant code is one 'word' (no spaces) with a letter in it, or a long number."""
    if v is None or isinstance(v, (bool, datetime, date)):
        return False
    s = str(v).strip()
    if not s or " " in s or len(s) > 30:
        return False
    n = norm(s)
    if n in NOT_CODES:
        return False
    return bool(re.search(r"[A-Z]", n)) or (n.isdigit() and len(n) >= 4)


def cell_codes(v):
    """Codes in one cell: the cell itself, or each part of "A1459741 + A1156216"."""
    if looks_like_code(v):
        return [norm(v)]
    if isinstance(v, str) and re.search(SPLIT, v):
        return [norm(p) for p in re.split(SPLIT, v) if looks_like_code(p) and re.search(r"\d", p)]
    return []


def day_key(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None


def text(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return day_key(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def money(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        # leading number, like "54/HR" -> 54 (same as the web app)
        m = re.match(r"[-+]?(\d+\.?\d*|\.\d+)", re.sub(r"[£,\s]", "", v))
        return float(m.group(0)) if m else None
    return None


def words(s):
    """Words worth comparing: 3+ letters, plurals folded ("ladders" = "ladder"). Keeps order."""
    out = {}
    for w in re.findall(r"[a-z]{3,}", str(s or "").lower()):
        if w in STOP:
            continue
        if len(w) > 4 and w.endswith("s") and not w.endswith("ss"):
            w = w[:-1]
        out[w] = True
    return list(out)


def measures(s):
    """Sizes: "4m", "350mm", "3.0 kva", "6/MTR" -> {"4m": "m", "350mm": "mm", "3kva": "kva", "6m": "m"}"""
    out = {}
    for num, unit in re.findall(r"(\d+(?:\.\d+)?)\s*/?\s*(mm|mtrs?|metres?|meters?|m|tonnes?|t|kva|cfm|ft)(?![a-z])", str(s or ""), re.I):
        unit = unit.lower()
        if re.match(r"(mtr|metre|meter)", unit):
            unit = "m"
        if unit.startswith("tonne"):
            unit = "t"
        f = float(num)
        out[(str(int(f)) if f.is_integer() else repr(f)) + unit] = unit
    return out


def qty_from(s):
    """Quantity written into a description: "2no", "2 no.", "x3", "3 x"."""
    s = s or ""
    m = re.search(r"(\d+)\s*no\b\.?", s, re.I) or re.search(r"\bx\s*(\d+)\b", s, re.I) or re.search(r"^\s*(\d+)\s*x\b", s, re.I)
    return int(m.group(1)) if m else 1


def gbp(n):
    return f"£{n:.2f}" if n else "£0"


def uk(d):
    return f"{d[8:10]}/{d[5:7]}/{d[2:4]}" if d else ""


def open_xlsx(path):
    """Old .xls files get converted to .xlsx first."""
    if path.lower().endswith(".xls"):
        out = tempfile.mkdtemp()
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, path],
                       capture_output=True, check=True)
        path = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
    return path


def read_inhouse(path):
    """Every code-like value on every sheet (same as Ctrl+F across the workbook), the cells it's in,
    and the plant lines under each header row."""
    wb = load_workbook(open_xlsx(path), data_only=True)
    codes, locs, items = {}, {}, []
    for ws in wb.worksheets:
        cols = None
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            for c, v in enumerate(row):
                for code in cell_codes(v):
                    codes.setdefault(code, set()).add(ws.title)
                    locs.setdefault(code, []).append(f"{ws.title}!{get_column_letter(c + 1)}{r}")
            # header row: names the company and asset columns (can repeat mid-sheet)
            hdr = [v.strip().lower() if isinstance(v, str) else "" for v in row]
            find = lambda pat: next((i for i, h in enumerate(hdr) if re.search(pat, h)), -1)
            asset, company = find(r"asset|plant\s*code|^code$"), find(r"company|supplier|hire\s*co")
            if asset >= 0 and company >= 0:
                cols = dict(asset=asset, company=company, width=max(i for i, h in enumerate(hdr) if h) + 1,
                            desc=find(r"detail|desc|plant hired"), date=find(r"date"),
                            rate=find(r"rate|price|cost"), status=find(r"status"))
                continue
            if not cols:
                continue
            get = lambda i: row[i] if 0 <= i < len(row) else None
            asset_cell = get(cols["asset"])
            parts = cell_codes(asset_cell)
            if (not parts and isinstance(asset_cell, str) and re.search(r"[A-Za-z]", asset_cell)
                    and re.search(r"\d", asset_cell) and len(asset_cell.strip()) <= 30):
                parts = [norm(asset_cell)]   # "P LADDERS 4M CQ": kept for suggestions, not added to the code list
            desc, comp = text(get(cols["desc"])), text(get(cols["company"]))
            if not parts or (not desc and not comp):
                continue                     # blank lines, colour keys
            rate = money(get(cols["rate"]))
            items.append(dict(sheet=ws.title, row=r, codes=parts, width=cols["width"],
                              cell=f"{ws.title}!{get_column_letter(cols['asset'] + 1)}{r}",
                              assetText=text(asset_cell), company=comp, desc=desc,
                              date=day_key(get(cols["date"])), rate=rate, zero=not rate, qty=qty_from(desc),
                              offHire=bool(re.search("off", text(get(cols["status"])), re.I))))
    return dict(codes=codes, locs=locs, items=items)


def find_code_column(ws, codes):
    """The plant-code column is whichever column has the most hits against the in-house report."""
    best_col, best_hits = None, 0
    for col in ws.iter_cols():
        hits = sum(1 for c in col if looks_like_code(c.value) and norm(c.value) in codes)
        if hits > best_hits:
            best_col, best_hits = col[0].column, hits
    if best_col:
        return best_col
    # nothing matched at all - fall back to a column headed like a code column
    for row in ws.iter_rows(max_row=30):
        for c in row:
            if isinstance(c.value, str) and re.search(r"code|plant|item|product", c.value, re.I):
                return c.column
    return None


def supplier_header(vws, col, first_row):
    """Header row above the first charge line: the nearest row whose code-column cell is text but not a code."""
    for r in range(first_row - 1, 0, -1):
        v = vws.cell(r, col).value
        if isinstance(v, str) and v.strip() and not looks_like_code(v):
            hdr = [(vws.cell(r, c).value or "").strip().lower() if isinstance(vws.cell(r, c).value, str) else ""
                   for c in range(1, vws.max_column + 1)]

            def pick(*pats):
                for p in pats:
                    i = next((i for i, h in enumerate(hdr) if re.search(p, h)), -1)
                    if i >= 0:
                        return i + 1
                return None
            return dict(row=r, desc=[i + 1 for i, h in enumerate(hdr) if "desc" in h],
                        rate=pick(r"^rate$", r"\brate\b", r"price|amount|cost|charge$"),
                        date=pick(r"start", r"on.?hire", r"^date$", r"date"),
                        qty=pick(r"^qty$", r"^quantity$", r"\bqty\b", r"quantity"))
    return None


def check_supplier(path, acs, code_col=None):
    codes, locs = acs["codes"], acs["locs"]
    src = open_xlsx(path)
    wb = load_workbook(src)                    # keeps formulas/formatting for the output
    vals = load_workbook(src, data_only=True)  # calculated values for reading
    matched, flagged, seen = [], [], set()
    name = os.path.basename(path)

    for ws in wb.worksheets:
        vws = vals[ws.title]
        col = column_index_from_string(code_col) if code_col else find_code_column(vws, codes)
        if not col:
            continue
        rows = [r for r in range(1, vws.max_row + 1) if looks_like_code(vws.cell(r, col).value)]
        if not rows:
            continue
        hdr = supplier_header(vws, col, rows[0])
        width = vws.max_column
        extra = width + 1                      # the added "ACS match" column
        if hdr:
            ws.cell(hdr["row"], extra).value = "ACS match"
            ws.cell(hdr["row"], extra).font = ws.cell(hdr["row"], extra).font.copy(bold=True)
        for r in rows:
            v = vws.cell(r, col).value
            code = norm(v)
            seen.add(code)
            cell = lambda c: vws.cell(r, c).value if c else None
            full = [text(vws.cell(r, c).value) for c in range(1, width + 1)]
            full = [x for x in full if x and x != text(v)]
            desc = " ".join(dict.fromkeys(x for x in (text(cell(c)) for c in hdr["desc"]) if x)) if hdr and hdr["desc"] else ""
            line = dict(supplier=name, sheet=ws.title, row=r, code=code, raw=text(v), desc=desc or " · ".join(full[:3]),
                        rate=cell(hdr["rate"]) if hdr else None, date=day_key(cell(hdr["date"])) if hdr else None,
                        qty=(money(cell(hdr["qty"])) or 1) if hdr else 1, full=full)
            line["charged"] = True if not hdr or not hdr["rate"] else bool(money(line["rate"]))
            if code in codes:
                line["acsAt"] = locs[code]
                ws.cell(r, extra).value = "; ".join(locs[code])
                for c in range(1, extra + 1):
                    target = ws.cell(r, c)
                    if isinstance(target, MergedCell):
                        continue
                    target.fill = GREEN
                    target.font = target.font.copy(bold=True)
                matched.append(line)
            else:
                # not on our report - left untouched in the file, reported on screen instead
                flagged.append(line)

    # Which supplier on the ACS report this is: the company behind most of the matches
    matched_codes = {m["code"] for m in matched}
    tally = {}
    for it in acs["items"]:
        if any(c in matched_codes for c in it["codes"]):
            k = it["company"].lower()
            tally[k] = tally.get(k, 0) + 1
    key = max(tally, key=tally.get) if tally else None
    company = next(i["company"] for i in acs["items"] if i["company"].lower() == key) if key else None
    for f in flagged:
        f["company"] = company

    base = os.path.splitext(os.path.basename(path))[0]
    out = os.path.join(os.path.dirname(os.path.abspath(path)), f"{base}_checked.xlsx")
    wb.save(out)
    return dict(name=name, out=out, matched=matched, flagged=flagged, company=company)


# ---------- suggestions: run once every supplier report has been checked ----------
def compare(f, a, use_rate):
    """How alike a supplier line and an ACS line are. None when they can't be the same plant."""
    score, reasons, date_eq, rate_eq = 0, [], False, False
    fm, am = measures(f["desc"]), measures(a["desc"])
    same_size = [k for k in fm if k in am]
    if {u for u in fm.values() if u in am.values()} and not same_size:
        return None                            # sizes disagree (4m vs 5m ladder)
    if same_size:
        score += 2 * min(len(same_size), 2)
        reasons.append("both " + ", ".join(same_size))
    aw = words(a["desc"])
    w = [x for x in words(f["desc"]) if x in aw]
    if w:
        score += min(len(w), 3)
        reasons.append('both "' + " ".join(w[:3]) + '"')
    if f["date"] and a["date"]:
        d = abs((date.fromisoformat(f["date"]) - date.fromisoformat(a["date"])).days)
        if d == 0:
            score += 3; reasons.append("same on-hire date"); date_eq = True
        elif d <= 3:
            score += 1; reasons.append(f"on-hire dates {d} day{'s' if d > 1 else ''} apart")
    fr, ar = money(f["rate"]), a["rate"]
    if use_rate and fr and fr > 0 and ar and ar > 0:
        if abs(fr - ar) < 0.005:
            score += 3; reasons.append(f"same rate {gbp(fr)}"); rate_eq = True
        elif a["qty"] > 1 and abs(fr * a["qty"] - ar) < 0.005:
            score += 2; reasons.append(f"{gbp(fr)} × {a['qty']} = ACS {gbp(ar)}"); rate_eq = True
    if f["qty"] > 1 and f["qty"] == a["qty"]:
        score += 1; reasons.append(f"both quantity {a['qty']}")
    if f.get("company") and a["company"] and f["company"].lower() == a["company"].lower():
        score += 1; reasons.append(f"ACS lists it under {a['company']}")
    return dict(score=score, reasons=reasons, date_eq=date_eq, rate_eq=rate_eq, words=len(w), sizes=len(same_size))


def likely_same(c):
    """Same plant, worded differently: the description agrees AND the date or rate agrees, or both date and rate agree."""
    return c["score"] >= 4 and (((c["words"] or c["sizes"]) and (c["date_eq"] or c["rate_eq"])) or (c["date_eq"] and c["rate_eq"]))


def likely_zero(c):
    """£0 lines have no rate to compare: the description must share a word (a ladder is a ladder)."""
    return c["words"] > 0 and c["score"] >= 2


def pair_up(lines, items, use_rate, accept):
    """Pairs highest scores first, each line used once (an ACS line once per unpaired asset on it)."""
    pairs = []
    for f in lines:
        for a in items:
            c = compare(f, a, use_rate)
            if c and accept(c):
                pairs.append((f, a, c))
    pairs.sort(key=lambda p: -p[2]["score"])
    used, out = set(), []
    for f, a, c in pairs:
        if id(f) in used or a["left"] <= 0:
            continue
        used.add(id(f)); a["left"] -= 1
        out.append((f, a, c))
    return out


def cross_check(acs, results):
    paired = {m["code"] for r in results for m in r["matched"]}
    open_items = []
    for a in acs["items"]:
        free = [c for c in a["codes"] if c not in paired]
        if free:
            open_items.append(dict(a, free=free, left=len(free)))
    unmatched = [f for r in results for f in r["flagged"]]

    # 1. Plant code not on the ACS report: best unpaired ACS line by date, rate, quantity and description
    found = pair_up(unmatched, open_items, True, likely_same)
    by_line = {id(f): (a, c) for f, a, c in found}
    suggestions = [(f, *by_line.get(id(f), (None, None))) for f in unmatched if f["charged"] or id(f) in by_line]

    # 2. End pass: £0 ACS lines still unpaired, looked for among the supplier lines nothing else claimed
    zero_open = [a for a in open_items if a["zero"] and a["left"] == len(a["free"])]
    spare = [f for f in unmatched if id(f) not in by_line]
    zero_found = {a["cell"]: (f, c) for f, a, c in pair_up(spare, zero_open, False, likely_zero)}
    zero_pass = [(a, *zero_found.get(a["cell"], (None, None))) for a in zero_open]
    return suggestions, zero_pass


def mark_inhouse(path, acs, results):
    """Copy of the ACS report: every plant line charged on a supplier report in bold with a thick border.
    The report's own colours are kept. A line listing several assets is only marked when all are charged."""
    charged = {m["code"] for r in results for m in r["matched"]}
    lines = [i for i in acs["items"] if all(c in charged for c in i["codes"])]
    wb = load_workbook(open_xlsx(path))
    for l in lines:
        ws = wb[l["sheet"]]
        for c in range(1, l["width"] + 1):
            cell = ws.cell(l["row"], c)
            if isinstance(cell, MergedCell):
                continue
            b = cell.border
            cell.border = Border(left=THICK if c == 1 else b.left, right=THICK if c == l["width"] else b.right,
                                 top=THICK, bottom=THICK)
            cell.font = cell.font.copy(bold=True)
    base = os.path.splitext(os.path.basename(path))[0]
    out = os.path.join(os.path.dirname(os.path.abspath(path)), f"{base}_checked.xlsx")
    wb.save(out)
    return out, len(lines)


def report_date(name):
    """ACS report date from its filename, e.g. "01.09.26"."""
    m = re.search(r"(\d{2})[.\-_](\d{2})[.\-_](\d{2,4})", name)
    if not m:
        return None
    y = m.group(3) if len(m.group(3)) == 4 else "20" + m.group(3)
    return f"{y}-{m.group(2)}-{m.group(1)}"


def sentence(parts):
    return ", ".join(parts[:-1]) + " and " + parts[-1] if len(parts) > 1 else (parts[0] if parts else "")


if __name__ == "__main__":
    args = sys.argv[1:]
    code_col = None
    if "--code-col" in args:
        i = args.index("--code-col")
        code_col = args[i + 1]
        args = args[:i] + args[i + 2:]
    if len(args) < 2:
        sys.exit("Usage: python3 hire_check.py ACS_REPORT.xlsx SUPPLIER1.xlsx [SUPPLIER2.xlsx ...] [--code-col C]")
    inhouse, suppliers = args[0], args[1:]

    acs = read_inhouse(inhouse)
    results = []
    all_flagged = 0
    for s in suppliers:
        res = check_supplier(s, acs, code_col)
        results.append(res)
        m, f = res["matched"], res["flagged"]
        all_flagged += len(f)
        if not f:
            print(f"OK  {res['name']}: all {len(m)} charged items are on the ACS report")
            continue
        print(f"!!  {res['name']}: {len(f)} charged item(s) NOT on the ACS report ({len(m)} matched)")
        for line in f:
            print(f"      row {line['row']}: {line['raw']}  -  {'  |  '.join(line['full'])}")
    print(f"\nTOTAL NOT ON ACS REPORT: {all_flagged}")

    suggestions, zero_pass = cross_check(acs, results)
    cut = report_date(os.path.basename(inhouse))
    many = len(results) > 1
    print("\nSUGGESTIONS - plant codes not on the ACS report")
    if not suggestions:
        print("  Every charged plant code is on the ACS report.")
    for f, a, c in suggestions:
        who = f"{f['supplier']}, " if many else ""
        print(f"  {who}row {f['row']}: {f['raw']} {f['desc']}, {gbp(money(f['rate']))}" + (f", on hire {uk(f['date'])}." if f["date"] else "."))
        if a:
            print(f"      Probably ACS {a['cell']}: {a['desc']}, asset {', '.join(a['free'])}, {gbp(a['rate'])}"
                  + (", marked off-hire on the ACS report" if a["offHire"] else "")
                  + f". Reason: {sentence(c['reasons'])}. If it's the same plant, update the asset number on the ACS report.")
        elif f["date"] and cut and f["date"] > cut:
            print("      No likely ACS line found. It went on hire after the ACS report date, so it may be new plant to add.")
        else:
            print("      No likely ACS line found. Check whether it should be on the ACS report.")

    print("\n£0 ITEMS ON THE ACS REPORT - checked last")
    if not zero_pass:
        print("  No unpaired £0 items on the ACS report.")
    checked = {(r["company"] or "").lower() for r in results}
    for a, f, c in zero_pass:
        head = f"  {a['cell']} {a['desc']} ({a['company'] or 'no company'}, £0):"
        if f:
            print(f"{head} found on {f['supplier']} row {f['row']}: {f['raw']} {f['desc']}, {gbp(money(f['rate']))}. Reason: {sentence(c['reasons'])}.")
        else:
            missing = a["company"] and a["company"].lower() not in checked
            print(f"{head} not found on the supplier reports checked." + (f" No {a['company']} report has been checked yet." if missing else ""))

    acs_out, acs_marked = mark_inhouse(inhouse, acs, results)
    print(f"\nACS report: {acs_marked} line(s) charged on the supplier reports in bold with a thick border -> {os.path.basename(acs_out)}")
